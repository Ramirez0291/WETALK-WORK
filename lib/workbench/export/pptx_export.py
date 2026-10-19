# 工坊导出（WP7）：幻灯片源 → PPTX（python-pptx 套用模板的版式），以及读回核对用的读取。
# 用法（输入输出都走 UTF-8 文件，不经控制台，路径含中文、日文时也不会乱码）：
#   pptx_export.py build    <spec.json> <out.pptx> <result.json>
#   pptx_export.py readback <in.pptx> <result.json>
#   pptx_export.py layouts  <template.pptx> <result.json>
# 只写命令行里给出的输出文件。spec：{ title, template, plan, slides:[{ id, version, title, paragraphs:[{ text, level, bullet }], notes }] }
import json
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
EMU_PER_PT = 12700


def pt(v):
    return round((v or 0) / EMU_PER_PT, 2)


PH_KIND = {
    PP_PLACEHOLDER.TITLE: 'title', PP_PLACEHOLDER.CENTER_TITLE: 'title', PP_PLACEHOLDER.VERTICAL_TITLE: 'title',
    PP_PLACEHOLDER.SUBTITLE: 'subtitle',
    PP_PLACEHOLDER.BODY: 'body', PP_PLACEHOLDER.VERTICAL_BODY: 'body',
    PP_PLACEHOLDER.OBJECT: 'object', PP_PLACEHOLDER.VERTICAL_OBJECT: 'object',
    PP_PLACEHOLDER.PICTURE: 'picture', PP_PLACEHOLDER.CHART: 'chart', PP_PLACEHOLDER.TABLE: 'table',
    PP_PLACEHOLDER.DATE: 'meta', PP_PLACEHOLDER.FOOTER: 'meta', PP_PLACEHOLDER.SLIDE_NUMBER: 'meta', PP_PLACEHOLDER.HEADER: 'meta',
}


def ph_kind(ph):
    try:
        return PH_KIND.get(ph.placeholder_format.type, 'other')
    except Exception:
        return 'other'


def layout_info(layout, index):
    return {
        'index': index,
        'name': layout.name,
        'placeholders': [{'idx': ph.placeholder_format.idx, 'type': ph_kind(ph), 'name': ph.name} for ph in layout.placeholders],
    }


def kinds_of(layout):
    return {ph_kind(ph) for ph in layout.placeholders}


def pick_layout(prs, want_body):
    named = ['Title and Content', '标题和内容', 'タイトルとコンテンツ'] if want_body else ['Title Only', '仅标题', 'タイトルのみ']
    for layout in prs.slide_layouts:
        if layout.name in named:
            return layout
    for layout in prs.slide_layouts:
        k = kinds_of(layout)
        if 'title' in k and (bool(k & {'body', 'object'}) == want_body):
            return layout
    return prs.slide_layouts[0]


def clear_slides(prs):
    # 模板里自带的示例页一律去掉，只留版式
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def set_bullet(para, bullet, level):
    p_pr = para._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in p_pr.findall(qn(tag)):
            p_pr.remove(el)
    if bullet:
        return  # 要点：沿用模板（或下面补上的）项目符号
    p_pr.set('marL', '0')
    p_pr.set('indent', '0')
    p_pr.append(p_pr.makeelement(qn('a:buNone'), {}))


def fill(frame, paragraphs, explicit_bullets):
    frame.clear()
    for i, p in enumerate(paragraphs):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = p['text']
        para.level = max(0, min(4, int(p.get('level', 0))))
        bullet = bool(p.get('bullet'))
        set_bullet(para, bullet, para.level)
        if bullet and explicit_bullets:
            # 自己加的文本框没有母版的项目符号：显式写一个
            p_pr = para._p.get_or_add_pPr()
            p_pr.set('marL', str(Emu(Pt(18 + 18 * para.level))))
            p_pr.set('indent', str(-Emu(Pt(18))))
            p_pr.append(p_pr.makeelement(qn('a:buChar'), {'char': '•'}))


def build(spec_path, out_path):
    with open(spec_path, encoding='utf-8') as f:
        spec = json.load(f)
    prs = Presentation(spec['template'])
    clear_slides(prs)
    by_name = {layout.name: layout for layout in prs.slide_layouts}
    plan = {x['id']: x['layout'] for x in ((spec.get('plan') or {}).get('slides') or [])}
    with_body = pick_layout(prs, True)
    title_only = pick_layout(prs, False)
    width, height = prs.slide_width, prs.slide_height
    for s in spec['slides']:
        paras = [p for p in s['paragraphs'] if p['text']]
        if s['id'] in plan:
            layout = by_name[plan[s['id']]]
        else:
            layout = with_body if paras else title_only
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is None:
            title = slide.shapes.add_textbox(int(width * 0.06), int(height * 0.05), int(width * 0.88), int(height * 0.15))
        title.text_frame.text = s['title']
        body = next((ph for ph in slide.placeholders if ph_kind(ph) in ('body', 'object')), None)
        if paras:
            if body is None:
                body = slide.shapes.add_textbox(int(width * 0.06), int(height * 0.24), int(width * 0.88), int(height * 0.68))
                body.text_frame.word_wrap = True
                fill(body.text_frame, paras, True)
            else:
                fill(body.text_frame, paras, False)
        # 没用上的空占位符去掉（不然放映时什么都没有，编辑时却提示“单击此处添加文本”）
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)
        slide.notes_slide.notes_text_frame.text = s['notes']
    prs.save(out_path)
    return {'slides': len(spec['slides']), 'layouts': [layout_info(l, i) for i, l in enumerate(prs.slide_layouts)]}


def theme_fonts(prs):
    out = {}
    try:
        part = prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        for which, key in (('majorFont', 'major'), ('minorFont', 'minor')):
            node = root.find(f'.//{A}{which}')
            if node is None:
                continue
            for script in ('latin', 'ea'):
                el = node.find(f'{A}{script}')
                if el is not None and el.get('typeface'):
                    out[f'+{key[:2]}-{"lt" if script == "latin" else "ea"}'] = el.get('typeface')
    except Exception:
        pass
    return out


def run_fonts(frame, theme):
    names = set()
    for el in frame._txBody.iter(f'{A}latin', f'{A}ea', f'{A}cs'):
        face = el.get('typeface')
        if face:
            names.add(theme.get(face, face) if face.startswith('+') else face)
    return names


def shape_info(sh, theme, fonts):
    info = {'name': sh.name, 'left': pt(sh.left), 'top': pt(sh.top), 'width': pt(sh.width), 'height': pt(sh.height)}
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        info['kind'] = 'group'
        info['children'] = [shape_info(c, theme, fonts) for c in sh.shapes]
        return info
    if getattr(sh, "has_chart", False):
        info['kind'] = 'chart'
        return info
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        info['kind'] = 'picture'
        try:
            w, h = sh.image.size
            vis_w = w * (1 - (sh.crop_left or 0) - (sh.crop_right or 0))
            vis_h = h * (1 - (sh.crop_top or 0) - (sh.crop_bottom or 0))
            info['image_ratio'] = round(vis_w / vis_h, 4) if vis_h else None
            info['shape_ratio'] = round(sh.width / sh.height, 4) if sh.height else None
        except Exception as e:
            info['image_error'] = str(e)[:200]
        return info
    if getattr(sh, 'has_table', False) and sh.has_table:
        info['kind'] = 'table'
        return info
    if sh.has_text_frame:
        info['kind'] = 'text'
        info['text'] = sh.text_frame.text
        info['placeholder'] = ph_kind(sh) if sh.is_placeholder else None
        fonts.update(run_fonts(sh.text_frame, theme))
        return info
    info['kind'] = 'other'
    return info


def readback(in_path):
    prs = Presentation(in_path)
    theme = theme_fonts(prs)
    out = {'slide_width': pt(prs.slide_width), 'slide_height': pt(prs.slide_height), 'theme_fonts': sorted(set(theme.values())), 'slides': []}
    for index, slide in enumerate(prs.slides, 1):
        fonts = set()
        title_shape = slide.shapes.title
        shapes = [shape_info(sh, theme, fonts) for sh in slide.shapes]
        body = []
        for sh in slide.shapes:
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.has_text_frame and not (sh.is_placeholder and ph_kind(sh) == 'meta'):
                body.append(sh.text_frame.text)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
        out['slides'].append({
            'index': index,
            'layout': slide.slide_layout.name,
            'title': title_shape.text_frame.text if title_shape is not None else '',
            'body': '\n'.join(body),
            'notes': notes,
            'shapes': shapes,
            'fonts': sorted(fonts),
        })
    return out


def layouts(template):
    prs = Presentation(template)
    return {'layouts': [layout_info(l, i) for i, l in enumerate(prs.slide_layouts)], 'theme_fonts': sorted(set(theme_fonts(prs).values()))}


def main():
    cmd = sys.argv[1]
    if cmd == 'build':
        result = build(sys.argv[2], sys.argv[3])
        dest = sys.argv[4]
    elif cmd == 'readback':
        result = readback(sys.argv[2])
        dest = sys.argv[3]
    elif cmd == 'layouts':
        result = layouts(sys.argv[2])
        dest = sys.argv[3]
    else:
        raise SystemExit(f'unknown command {cmd}')
    with open(dest, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    main()
